"""Shared pipeline used by the UI: PDFs in, compiled Admin/Absence/Overtime report out."""

import contextlib
import io
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

from mapper import load_templates, process_timesheet
from pdf2excel import convert as pdf_to_excel

# When packaged with PyInstaller the templates are unpacked next to the code in _MEIPASS
_BASE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
TEMPLATES_DIR = _BASE_DIR / "templates"


def build_report(files, expected_month=None):
    """files: list of (file_name, pdf_bytes). Returns (report, results, extracted).

    report is {"Admin": df, "Absence": df, "Overtime": df}; results is one dict per
    file with its name, status ("OK" / "Rejected") and a message. A file that fails
    is skipped and never contributes rows to the report.

    extracted is {file_name: rows}: the raw PDF -> Excel extraction of every file whose
    PDF could be read (also the rejected ones), as lists of cell values, so the user
    can check what was read from each PDF before it was mapped.
    """
    df_admin, df_absence, df_overtime = load_templates(TEMPLATES_DIR)
    results = []
    extracted = {}

    with tempfile.TemporaryDirectory() as tmp:
        for i, (name, data) in enumerate(files):
            try:
                pdf_path = Path(tmp) / f"{i}.pdf"
                pdf_path.write_bytes(data)
                # A windowed .exe has no console, so keep library output off stdout/stderr
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    xlsx_path = pdf_to_excel(pdf_path)
                extracted[name] = _sheet_rows(xlsx_path)
                # Map into copies so a failure half-way never leaves partial rows behind
                df_admin, df_absence, df_overtime = process_timesheet(
                    xlsx_path, df_admin, df_absence, df_overtime, expected_month
                )
                results.append({"File": name, "Status": "OK", "Message": ""})
            except (Exception, SystemExit) as e:
                results.append({"File": name, "Status": "Rejected", "Message": str(e)})

    report = {"Admin": df_admin, "Absence": df_absence, "Overtime": df_overtime}
    return report, results, extracted


def _sheet_rows(xlsx_path):
    """All cell values of the first sheet as a list of rows (trailing empty rows dropped)."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def extracted_to_xlsx_bytes(extracted):
    """One sheet per PDF holding exactly what was extracted from it (before mapping)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    for name, rows in extracted.items():
        base = re.sub(r"[\\/*?:\[\]]", "_", Path(name).stem)[:28] or "Sheet"
        title, n = base, 1
        while title.casefold() in used:
            n += 1
            title = f"{base[:28 - len(str(n))]}_{n}"
        used.add(title.casefold())
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    if not extracted:
        wb.create_sheet("No data")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_to_xlsx_bytes(report):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, df in report.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    return buffer.getvalue()
