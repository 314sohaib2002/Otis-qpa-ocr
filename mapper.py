import math

import openpyxl
import pandas as pd

# Layout of the monthly report (the PDF/Excel has one row per day, so everything
# below the day rows moves with the length of the month). Rows are located from
# the "SUM denne måned" label in column A, never from fixed row numbers:
#   rows 6 .. SUM-1  one row per day (28-31 rows)
#   SUM              monthly totals (E hours, F deviation, H 50%, I 100%, K vacation days)
#   SUM+1            total overtime this month
#   SUM+2            overtime hours to time off
#   SUM+3            overtime hours to pay out
#   SUM+4            time balance carried to next month (E)
#   SUM+5            remaining holiday (K)
FIRST_DAY_ROW = 6
SUM_LABEL = "sum denne"

# Row (1-based) holding the real column headers in each template workbook; the rows
# above it are titles/instructions and the rows below it are example/legend text.
TEMPLATE_HEADER_ROWS = {"Admin": 1, "Absence": 4, "Overtime": 8}

# Month labels as they appear in the report (Norwegian, plus English) -> month number
MONTH_ALIASES = {
    1: ("jan",), 2: ("feb",), 3: ("mar",), 4: ("apr",), 5: ("mai", "may"), 6: ("jun",),
    7: ("jul",), 8: ("aug",), 9: ("sep",), 10: ("okt", "oct"), 11: ("nov",), 12: ("des", "dec"),
}
MONTH_NAMES = {1: "January", 2: "February", 3: "March", 4: "April", 5: "May", 6: "June",
               7: "July", 8: "August", 9: "September", 10: "October", 11: "November",
               12: "December"}


def month_number(label):
    """Month number (1-12) for a label such as 'Jun' or 'Mai', or None."""
    text = str(label or "").strip().casefold()[:3]
    return next((n for n, names in MONTH_ALIASES.items() if text in names), None)


# Standard wage code names, as listed in the legend of the Overtime template
WAGE_CODE_NAMES = {
    118: "Overtime hours for pay out",
    425: "Overtime add 50%",
    426: "Overtime add 100%",
    61: "Overtime hours for time off",
    62: "Withdrawal of time off",
    52: "Timebank payout",
}


def clean_val(val):
    if val is None or val == "-":
        return ""
    return str(val)


def to_number(val):
    """Numeric value of a cell or None. Accepts numbers and text such as
    '32.00', '4,5' or '1.016,72' (PDF exports can deliver numbers as text)."""
    if isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        number = float(val)
    elif isinstance(val, str):
        text = val.strip()
        if "," in text:  # European format: dots are thousands separators
            text = text.replace(".", "").replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None
    else:
        return None
    return number if math.isfinite(number) else None


def num_str(val):
    """Cell value as text for the output sheets; numbers lose a needless '.0'."""
    number = to_number(val)
    if number is None:
        return clean_val(val)
    return str(int(number)) if number.is_integer() else str(round(number, 2))


def _empty_template(path, header_row):
    """Empty DataFrame with only the template's real header row as columns."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        header = next(wb.active.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True))
    finally:
        wb.close()
    names = ["" if h is None else str(h) for h in header]
    # The table ends at the first run of 2+ empty headers after its first header;
    # anything further right is a side legend (e.g. the wage code list in the
    # Overtime template). A single empty header inside the table is kept.
    first = next((i for i, n in enumerate(names) if n), 0)
    for i in range(first, len(names) - 1):
        if names[i] == "" and names[i + 1] == "":
            names = names[:i]
            break
    while names and names[-1] == "":
        names.pop()
    return pd.DataFrame(columns=names)


def load_templates(base_path):
    """Load the Admin, Absence and Overtime templates as empty DataFrames.

    Only the header row is kept, so the instruction/legend text in the templates
    never ends up as data rows in the compiled report.
    """
    return tuple(
        _empty_template(base_path / f"{file}.xlsx", TEMPLATE_HEADER_ROWS[name])
        for name, file in (("Admin", "Admin"), ("Absence", "Absense"), ("Overtime", "overtime"))
    )


def find_sum_row(sheet):
    """Row of the monthly totals, found by its column-A label."""
    for row in range(FIRST_DAY_ROW, sheet.max_row + 1):
        label = sheet[f"A{row}"].value
        if isinstance(label, str) and label.strip().casefold().startswith(SUM_LABEL):
            day_rows = row - FIRST_DAY_ROW
            if not 28 <= day_rows <= 31:
                raise ValueError(
                    f"Found the SUM row at {row}, which implies {day_rows} day rows "
                    "(expected 28-31). The layout may have changed."
                )
            return row
    raise ValueError(
        f"Could not find the '{SUM_LABEL.upper()} ...' row in column A. "
        "The layout may have changed - refusing to guess cell positions."
    )


def process_timesheet(input_excel_path, df_admin, df_absence, df_overtime, expected_month=None):
    # Load the workbook (data_only=True ensures we get the value, not formulas)
    wb = openpyxl.load_workbook(input_excel_path, data_only=True)
    sheet = wb.active

    sum_row = find_sum_row(sheet)

    if expected_month is not None:
        sheet_month = month_number(sheet["I2"].value)
        if sheet_month != expected_month:
            raise ValueError(
                f"This timesheet is for '{clean_val(sheet['I2'].value) or 'unknown month'}', "
                f"not {MONTH_NAMES[expected_month]}."
            )

    def at(col, offset=0):
        """Raw value of a cell in the totals block, relative to the SUM row."""
        return sheet[f"{col}{sum_row + offset}"].value

    # ==========================================
    # 1. PROCESS ADMIN TEMPLATE
    # ==========================================
    employee_id = clean_val(sheet["F2"].value)
    employee_name = clean_val(sheet["B2"].value)
    if not to_number(sheet["F2"].value) and not employee_id.strip("0 "):
        raise ValueError("Employee number (F2) is empty - cannot tell whose timesheet this is.")

    # "Sum overtid denne måneden" (the row below the SUM row)
    total_overtime = num_str(at("E", 1))

    # Payout = column G of the SUM row, or 50% + 100% overtime (H + I) when G is empty
    if to_number(at("G")) is None:
        overtime_payout = num_str((to_number(at("H")) or 0) + (to_number(at("I")) or 0))
    else:
        overtime_payout = num_str(at("G"))

    overtime_50 = num_str(at("H"))
    overtime_100 = num_str(at("I"))

    # Time balance transferred
    time_balance = num_str(at("E", 4))

    # Vacation days and remaining
    days_taken = num_str(at("K"))
    remaining_holiday = num_str(at("K", 5))

    # Timebank logic
    e5_val = to_number(sheet["E5"].value) or 0
    e40_val = to_number(at("E", 4)) or 0

    timebank_math_result = num_str(e5_val - e40_val)
    timebank_payout = timebank_math_result
    time_off_withdrawal = timebank_math_result

    overtime_for_time_off = num_str(at("E", 2))

    # Advanced Parsing for Dates (one row per day, up to the SUM row)
    vacation_dates = []
    self_report_dates = []
    sick_leave_dates = []

    for row in range(FIRST_DAY_ROW, sum_row):
        day_val = to_number(sheet[f"A{row}"].value)
        if day_val is None or not day_val.is_integer():
            continue

        day_int = int(day_val)

        # Check Vacation (Column K)
        k_val = to_number(sheet[f"K{row}"].value)
        if k_val is not None and k_val > 0:
            vacation_dates.append(str(day_int))

        # Check Sick Leave / Absence (Column J) and Comments (Column D)
        j_val = to_number(sheet[f"J{row}"].value)
        d_val = sheet[f"D{row}"].value

        if j_val is not None and j_val > 0:
            if d_val and str(d_val).strip() != "":
                sick_leave_dates.append(str(day_int))
            else:
                self_report_dates.append(str(day_int))

    vacation_dates_str = ", ".join(vacation_dates)
    self_report_dates_str = ", ".join(self_report_dates)
    sick_leave_dates_str = ", ".join(sick_leave_dates)

    cols = df_admin.columns
    
    new_admin_row = {
        cols[0]: employee_id,
        cols[1]: employee_name,
        cols[5]: total_overtime,
        cols[6]: overtime_payout,
        cols[7]: overtime_50,
        cols[8]: overtime_100,
        cols[9]: overtime_for_time_off,
        cols[10]: time_off_withdrawal,
        cols[11]: timebank_payout,
        cols[12]: time_balance,
        cols[13]: vacation_dates_str,
        cols[14]: days_taken,
        cols[15]: remaining_holiday,
        cols[16]: self_report_dates_str,
        cols[17]: sick_leave_dates_str
    }
    df_admin = pd.concat([df_admin, pd.DataFrame([new_admin_row])], ignore_index=True)


    # ==========================================
    # 2. PROCESS ABSENCE TEMPLATE
    # ==========================================
    def group_consecutive_days(days_list):
        if not days_list:
            return []
        days_list = sorted([int(d) for d in days_list])
        groups = []
        start_day = days_list[0]
        prev_day = days_list[0]
        count = 1
        for day in days_list[1:]:
            if day == prev_day + 1:
                count += 1
                prev_day = day
            else:
                groups.append({"from": start_day, "to": prev_day, "days": count})
                start_day = day
                prev_day = day
                count = 1
        groups.append({"from": start_day, "to": prev_day, "days": count})
        return groups

    vacation_groups = group_consecutive_days(vacation_dates)
    self_report_groups = group_consecutive_days(self_report_dates)
    cert_sick_groups = group_consecutive_days(sick_leave_dates)

    absence_rows = []
    month_str = clean_val(sheet["I2"].value) # e.g. "Jun"
    # The year sits in K2 in the PDF export (J2 holds the "ÅR:" label); fall back to J2
    # for sheets where the label is merged away. Only accept a real 4-digit year.
    year_str = ""
    for addr in ("K2", "J2"):
        candidate = clean_val(sheet[addr].value)
        if candidate.isdigit() and len(candidate) == 4:
            year_str = candidate
            break

    def format_date(day):
        return f"{day}-{month_str}-{year_str}"

    for grp in vacation_groups:
        row_data = [None] * len(df_absence.columns)
        row_data[1] = employee_id
        row_data[2] = employee_name
        row_data[4] = "Ferie (Vacation)"
        row_data[5] = "100%"
        row_data[6] = format_date(grp["from"])
        row_data[7] = format_date(grp["to"])
        row_data[8] = grp["days"]
        absence_rows.append(row_data)

    for grp in self_report_groups:
        row_data = [None] * len(df_absence.columns)
        row_data[1] = employee_id
        row_data[2] = employee_name
        row_data[4] = "2"
        row_data[5] = "100%"
        row_data[6] = format_date(grp["from"])
        row_data[7] = format_date(grp["to"])
        row_data[8] = grp["days"]
        absence_rows.append(row_data)

    for grp in cert_sick_groups:
        row_data = [None] * len(df_absence.columns)
        row_data[1] = employee_id
        row_data[2] = employee_name
        row_data[4] = "Syk med legeerklæring"
        row_data[5] = "100%"
        row_data[6] = format_date(grp["from"])
        row_data[7] = format_date(grp["to"])
        row_data[8] = grp["days"]
        absence_rows.append(row_data)

    if absence_rows:
        new_abs_df = pd.DataFrame(absence_rows, columns=df_absence.columns)
        df_absence = pd.concat([df_absence, new_abs_df], ignore_index=True)


    # ==========================================
    # 3. PROCESS OVERTIME TEMPLATE
    # ==========================================
    wage_entries = [
        (118, to_number(overtime_payout) or 0),
        (425, to_number(overtime_50) or 0),
        (426, to_number(overtime_100) or 0),
        (61,  to_number(overtime_for_time_off) or 0),
        (62,  to_number(time_off_withdrawal) or 0),
        (52,  to_number(timebank_payout) or 0),
    ]

    overtime_rows = []
    columns_length = len(df_overtime.columns)

    for code, value in wage_entries:
        if value != 0:  # Only add codes the employee actually has hours for
            row_data = [None] * columns_length
            row_data[0] = employee_id
            row_data[1] = employee_name
            row_data[2] = code
            row_data[3] = WAGE_CODE_NAMES[code]  # Wage code name (standard)
            row_data[5] = int(value) if float(value).is_integer() else value  # Number of
            overtime_rows.append(row_data)

    if overtime_rows:
        new_ovt_df = pd.DataFrame(overtime_rows, columns=df_overtime.columns)
        df_overtime = pd.concat([df_overtime, new_ovt_df], ignore_index=True)

    return df_admin, df_absence, df_overtime
